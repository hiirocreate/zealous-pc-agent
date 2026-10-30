"""チャットに添付されたファイル・フォルダを読み取り、AIが分析できる形(文と表の要約)にする。
読むだけで、ファイルを書き換えたり外へ送ったりはしない。"""
import os
import re

CODE = (".py", ".js", ".ts", ".tsx", ".jsx", ".html", ".htm", ".css", ".scss", ".java", ".kt", ".c", ".h", ".cpp",
        ".hpp", ".cs", ".go", ".rs", ".php", ".rb", ".swift", ".sh", ".bat", ".ps1", ".sql", ".vue", ".svelte", ".r",
        ".lua", ".pl", ".dart", ".m")
TEXT = (".txt", ".md", ".log", ".json", ".xml", ".yaml", ".yml", ".toml", ".ini", ".cfg", ".conf", ".env.example",
        ".gitignore", ".dockerfile", ".csv", ".tsv") + CODE
OFFICE = (".docx", ".pptx", ".xlsx", ".pdf")
IMAGES = (".png", ".jpg", ".jpeg", ".gif", ".bmp", ".webp", ".tif", ".tiff", ".heic")
SKIP_DIRS = {".git", "node_modules", "__pycache__", ".venv", "venv", "dist", "build", ".idea", ".vscode", "$recycle.bin"}


def label(path):
    """チップに出す名前(フォルダは末尾に /)。"""
    name = os.path.basename(path.rstrip("\\/")) or path
    return name + ("/" if os.path.isdir(path) else "")


def _size(n):
    for unit in ("B", "KB", "MB", "GB"):
        if n < 1024 or unit == "GB":
            return f"{n:.0f}{unit}" if unit == "B" else f"{n:.1f}{unit}"
        n /= 1024


def _table_summary(rows, name):
    """表(行の一覧)の要約: 行数・列・先頭の行・数の列の合計/平均/最小/最大。"""
    import officedoc
    rows = [r for r in rows if any(str(c).strip() for c in r)]
    if not rows:
        return f"{name}: 空の表"
    head, body = [str(c) for c in rows[0]], rows[1:]
    out = [f"{name}: {len(body)}行 × {len(head)}列", "列: " + " | ".join(head)]
    for c, title in enumerate(head):
        vals = []
        for r in body:
            v = r[c] if c < len(r) else None
            if isinstance(v, (int, float)) and not isinstance(v, bool):
                vals.append(float(v))
            elif v is not None:
                num = officedoc._number(str(v))
                if num:
                    vals.append(float(num[0]))
        if vals and len(vals) >= max(1, len(body) // 2):
            out.append(f"  {title}: 合計 {sum(vals):,.6g} / 平均 {sum(vals) / len(vals):,.6g} / "
                       f"最小 {min(vals):,.6g} / 最大 {max(vals):,.6g}")
    shown = body[:25]
    out.append("先頭の行:")
    out += [" | ".join("" if v is None else str(v) for v in r) for r in shown]
    if len(body) > len(shown):
        out.append(f"…(残り {len(body) - len(shown)} 行)")
    return "\n".join(out)


def read_file(path, limit=12000):
    """1つのファイルの中身(AIに渡す文)。"""
    import tools
    ext = os.path.splitext(path)[1].lower()
    name = os.path.basename(path)
    try:
        size = os.path.getsize(path)
    except OSError:
        return f"{name}: 読み取れませんでした"
    head = f"【ファイル: {name}】({_size(size)}・{path})"
    try:
        if ext in (".csv", ".tsv"):
            import csv
            text = tools.file_text(path, limit=3_000_000)
            rows = list(csv.reader(text.splitlines(), delimiter="\t" if ext == ".tsv" else ","))
            return head + "\n" + _table_summary(rows, name)[:limit]
        if ext == ".xlsx":
            from openpyxl import load_workbook
            wb = load_workbook(path, read_only=True, data_only=True)
            parts = [f"シート: {', '.join(wb.sheetnames)}"]
            for ws in wb.worksheets[:6]:
                rows = [list(r) for _, r in zip(range(2000), ws.iter_rows(values_only=True))]
                parts.append(_table_summary(rows, f"シート「{ws.title}」"))
            return head + "\n" + "\n\n".join(parts)[:limit]
        if ext == ".pdf":
            from pypdf import PdfReader
            reader = PdfReader(path)
            text = "\n".join((p.extract_text() or "") for p in reader.pages[:60])
            more = f"\n…(全{len(reader.pages)}ページのうち最初の60ページ)" if len(reader.pages) > 60 else ""
            return head + f"\n{len(reader.pages)}ページ\n" + text[:limit] + more
        if ext in IMAGES:
            from PIL import Image
            with Image.open(path) as im:
                info = f"画像 {im.width}×{im.height} ({im.format})"
                exif = im.getexif() if hasattr(im, "getexif") else {}
                taken = exif.get(36867) or exif.get(306)
            return head + f"\n{info}" + (f"・撮影日時 {taken}" if taken else "") + \
                "\n(画像の中身は、文字の情報としては読み取れません)"
        if ext in OFFICE or ext in TEXT or size < 200_000:
            text = tools.file_text(path, limit=600_000)
            if ext not in OFFICE and "\x00" in text[:2000]:
                return head + "\n(文字ではないファイルのため、中身は読み取れません)"
            text = re.sub(r"[ \t]+\n", "\n", text)
            cut = f"\n…(長いので最初の{limit}文字だけ。全体は{len(text)}文字)" if len(text) > limit else ""
            return head + "\n" + text[:limit] + cut
        return head + "\n(この種類のファイルの中身は読み取れません)"
    except Exception as e:
        return head + f"\n(読み取りに失敗しました: {type(e).__name__})"


def read_folder(path, limit=16000):
    """フォルダ: 中身の一覧(ツリー)と、種類・大きさの集計、主なファイルの中身。"""
    entries, kinds, total = [], {}, 0
    base = path.rstrip("\\/")
    for cur, dirs, files in os.walk(base):
        dirs[:] = sorted(d for d in dirs if d.lower() not in SKIP_DIRS and not d.startswith("."))
        depth = cur[len(base):].count(os.sep)
        if depth >= 5:
            dirs[:] = []
            continue
        if cur != base:
            entries.append("  " * (depth - 1) + os.path.basename(cur) + "/")
        for f in sorted(files):
            p = os.path.join(cur, f)
            try:
                n = os.path.getsize(p)
            except OSError:
                continue
            total += n
            ext = os.path.splitext(f)[1].lower() or "(拡張子なし)"
            kinds[ext] = kinds.get(ext, 0) + 1
            entries.append("  " * depth + f"{f} ({_size(n)})")
        if len(entries) > 400:
            break
    files = sum(kinds.values())
    top = "、".join(f"{k} {v}件" for k, v in sorted(kinds.items(), key=lambda kv: -kv[1])[:10])
    out = [f"【フォルダ: {label(path)}】({path})", f"ファイル {files}件・合計 {_size(total)}・種類: {top}", "中身:"]
    out += entries[:250] + ([f"…(ほか {len(entries) - 250} 件)"] if len(entries) > 250 else [])
    text = "\n".join(out)
    # 主なファイル(説明書き・設定・小さい文書やコード)を、残りの量の分だけ読む
    order = []
    for cur, dirs, fs in os.walk(base):
        dirs[:] = [d for d in dirs if d.lower() not in SKIP_DIRS and not d.startswith(".")]
        for f in fs:
            ext = os.path.splitext(f)[1].lower()
            if ext in TEXT or ext in OFFICE:
                rank = 0 if re.match(r"readme|概要|説明|仕様", f, re.I) else 1 if ext in (".md", ".txt", ".docx", ".pdf") \
                    else 2 if ext in (".json", ".toml", ".yaml", ".yml", ".csv", ".xlsx") else 3
                order.append((rank, os.path.join(cur, f)))
    left = limit - len(text)
    for _, p in sorted(order)[:40]:
        if left < 800:
            break
        part = read_file(p, limit=min(4000, left - 200))
        text += "\n\n" + part
        left -= len(part)
    return text


def material(paths, limit=18000):
    """添付すべての中身(AIに渡す文)。量が多いときは、1つずつの分を減らす。"""
    paths = [p for p in paths if os.path.exists(p)]
    if not paths:
        return ""
    each = max(3000, limit // len(paths))
    parts = [read_folder(p, each) if os.path.isdir(p) else read_file(p, each) for p in paths]
    return "\n\n".join(parts)[:limit + 2000]
