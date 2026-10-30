import os
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import config  # noqa: E402
import llm  # noqa: E402
import tools  # noqa: E402
from wakeword import strip_wake_word, voice_command  # noqa: E402


def test_wake_word():
    assert strip_wake_word("ウェイク、ゼラス、メモ帳を作って") == "メモ帳を作って"
    assert strip_wake_word("Wake Zealous open youtube") == "open youtube"
    assert strip_wake_word("ゼラス、音量を50に") is None
    assert strip_wake_word("今日の天気は？") is None


def test_rename_ai(tmp_path, monkeypatch):
    monkeypatch.setattr(config, "SETTINGS_FILE", str(tmp_path / "s.json"))
    monkeypatch.setattr(config, "APP_DIR", str(tmp_path))
    config.save({"AI_NAME": "アトラス", "AI_NAME_ALIASES": []})
    assert strip_wake_word("ウェイク、アトラス、検索して") == "検索して"
    assert strip_wake_word("ウェイク、ゼラス、検索して") is None
    config.reset()
    config.load()
    assert config.AI_NAME == "Zealous"


def test_parse_tool_calls():
    msg = {"tool_calls": [{"function": {"name": "web_search", "arguments": {"query": "猫"}}}]}
    assert llm.parse_tool_calls(msg) == [("web_search", {"query": "猫"})]
    msg = {"content": '{"name": "lock_screen", "arguments": {}}'}
    assert llm.parse_tool_calls(msg) == [("lock_screen", {})]
    assert llm.parse_tool_calls({"content": "こんにちは"}) == []


def test_run_tool(tmp_path, monkeypatch):
    monkeypatch.setattr(config, "DESKTOP", str(tmp_path))
    tools.run_tool("create_text_file", {"name": "memo", "content": "abc", "bogus": 1})
    assert (tmp_path / "memo.txt").read_text(encoding="utf-8") == "abc"
    assert "対応していません" in tools.run_tool("nope", {})
    monkeypatch.setattr(config, "DRY_RUN", True)
    assert tools.run_tool("lock_screen", {}).startswith("[DRY]")


def test_ask_with_mock(monkeypatch):
    monkeypatch.setattr(config, "DRY_RUN", True)
    monkeypatch.setattr(llm, "_chat", lambda m, use_tools=True, **kw: {
        "tool_calls": [{"function": {"name": "open_url", "arguments": {"url": "example.com"}}}]})
    assert llm.ask("example.comを開いて").startswith("[DRY] open_url")


def test_folder_ops(tmp_path, monkeypatch):
    monkeypatch.setattr(config, "DESKTOP", str(tmp_path))
    tools.run_tool("create_folder", {"name": "A"})
    tools.run_tool("create_text_file", {"name": "m", "content": "x"})
    assert "[A]" in tools.run_tool("list_folder", {"path": str(tmp_path)})
    tools.run_tool("rename_item", {"name": "m.txt", "new_name": "n.txt"})
    tools.run_tool("move_item", {"name": "n.txt", "dest": str(tmp_path / "A")})
    assert (tmp_path / "A" / "n.txt").exists()
    tools.run_tool("copy_item", {"name": "A", "dest": str(tmp_path / "B")})
    assert (tmp_path / "B" / "A" / "n.txt").exists()
    # ホームフォルダそのものや設定で禁止された操作は拒否
    assert "できませんでした" in tools.run_tool("rename_item", {"name": os.path.expanduser("~"), "new_name": "x"})
    monkeypatch.setattr(config, "FOLDER_OPS_ENABLED", False)
    assert "無効" in tools.run_tool("create_folder", {"name": "C"})


def test_version_compare():
    import updater
    assert updater._ver("v1.10.0") > updater._ver("1.9.9")
    assert updater._ver("1.0") == updater._ver("v1.0.0")


def test_voice_models(tmp_path, monkeypatch):
    import json
    import voicevox
    cat = tmp_path / "cat.json"
    cat.write_text(json.dumps([{"id": 3, "speaker": "ずんだもん", "style": "ノーマル", "vvm": "0.vvm", "mb": 56, "uuid": "z"},
                               {"id": 8, "speaker": "春日部つむぎ", "style": "ノーマル", "vvm": "0.vvm", "mb": 56, "uuid": "k"},
                               {"id": 13, "speaker": "青山龍星", "style": "ノーマル", "vvm": "15.vvm", "mb": 63, "uuid": "a"}]),
                   encoding="utf-8")
    monkeypatch.setattr(voicevox, "CATALOG", str(cat))
    monkeypatch.setattr(voicevox, "BUNDLED", str(cat))
    monkeypatch.setattr(voicevox, "MODEL_DIR", str(tmp_path / "model"))
    assert voicevox.needs_change(3)                     # まだ何も入っていない
    (tmp_path / "model").mkdir()
    (tmp_path / "model" / "0.vvm").write_bytes(b"")
    assert not voicevox.needs_change(8)                 # 同じモデルに入っている声
    assert voicevox.needs_change(13)                    # 別のモデルが必要
    assert voicevox.install_voice(8, lambda *a: None) is False
    assert [(sp["name"], [st["id"] for st in sp["styles"]]) for sp in voicevox.speakers()] == [
        ("ずんだもん", [3]), ("春日部つむぎ", [8]), ("青山龍星", [13])]


def test_settings_fields():
    import pytest
    sw = pytest.importorskip("settings_window")  # 画面部品(tkinter)が無い環境では飛ばす
    assert {f[0] for f in sw.FIELDS} <= set(config.DEFAULTS)
    assert sw.parse("list", "a、b, ,c") == ["a", "b", "c"]


def test_voice_command():
    assert voice_command("メモ帳を開いて", True) == "メモ帳を開いて"            # ボタン: 合図不要
    assert voice_command("ウェイク、Zealous、メモ帳を開いて", True) == "メモ帳を開いて"
    assert voice_command("メモ帳を開いて", False) is None                      # 常時待ち受け: 合図が必要
    assert voice_command("ウェイク、Zealous、メモ帳を開いて", False) == "メモ帳を開いて"


def test_search_summary(monkeypatch):
    import websearch
    monkeypatch.setattr(websearch, "search", lambda q, limit=5: [{"title": "A", "url": "https://a", "snippet": "a"},
                                                        {"title": "B", "url": "https://b", "snippet": "b"}])
    monkeypatch.setattr(websearch, "read_pages", lambda r, count=3, limit=1500: ["本文A", "本文B"])
    monkeypatch.setattr(websearch, "images", lambda q, limit=30: [{"thumb": "t"}])
    monkeypatch.setattr(llm, "_chat", lambda m, use_tools=True, **kw: {"content": "要約です"})
    body, sources, images = llm.split_reply(tools.run_tool("search_summary", {"query": "猫"}))
    assert body == "要約です" and sources == "[1] A\nhttps://a\n[2] B\nhttps://b"
    assert images == {"query": "猫", "items": [{"thumb": "t"}], "open": False}  # 関連画像は一覧を自動で開かない


def test_parse_bing():
    import base64
    import websearch
    url = "https://example.com/page"
    enc = base64.urlsafe_b64encode(url.encode()).decode().rstrip("=")
    page = (f'<li class="b_algo"><h2><a href="https://www.bing.com/ck/a?!&amp;u=a1{enc}&amp;ntb=1">'
            '<strong>例</strong>のページ</a></h2><p>説明&amp;文</p></li>')
    assert websearch.parse_bing(page) == [{"title": "例のページ", "url": url, "snippet": "説明&文"}]


def test_find_app():
    import apps
    found = [{"Name": "Excel", "AppID": "1"}, {"Name": "Excel 2016 Database Compare", "AppID": "2"},
             {"Name": "Microsoft Edge", "AppID": "3"}]
    assert apps.find("エクセル", found)["AppID"] == "1"
    assert apps.find("edge", found)["AppID"] == "3"
    assert apps.find("存在しないアプリ", found) is None


def test_conversation_memory(monkeypatch):
    seen = []
    monkeypatch.setattr(config, "CONTEXT_TURNS", 1)
    monkeypatch.setattr(llm, "_chat", lambda m, use_tools=True, **kw: (seen.append(m), {"content": "はい"})[1])
    llm._history.clear()
    llm.ask("一つ目")
    llm.ask("二つ目")
    llm.ask("三つ目")
    contents = [m["content"] for m in seen[-1]]
    assert "二つ目" in contents and "一つ目" not in contents   # 直前の1往復だけ覚えている


def test_image_search(monkeypatch):
    import websearch
    items = [{"thumb": "t", "mid": "m", "full": "f", "title": "りんご", "page": "p"}]
    monkeypatch.setattr(websearch, "images", lambda q, limit=30: items)
    body, sources, images = llm.split_reply(tools.run_tool("image_search", {"query": "リンゴ"}))
    assert "1枚" in body and sources == ""
    assert images == {"query": "リンゴ", "items": items, "open": True}


def test_parse_yahoo_images():
    import json
    import websearch
    data = {"props": {"initialProps": {"pageProps": {"algos": [
        {"title": "赤いリンゴ", "refererUrl": "https://ex.jp", "thumbnail": {"url": "https://t"},
         "main": {"url": "https://m"}, "original": {"url": "https://o"}}]}}}}
    page = f'<script id="__NEXT_DATA__" type="application/json">{json.dumps(data)}</script>'
    assert websearch.parse_yahoo_images(page) == [
        {"thumb": "https://t", "mid": "https://m", "full": "https://o", "title": "赤いリンゴ", "page": "https://ex.jp"}]


def test_model_labels():
    import models
    text = models.label("qwen2.5:3b")
    assert "1.9GB" in text and models.name_of(text) == "qwen2.5:3b"
    assert models.name_of("gemma2:2b") == "gemma2:2b"          # 手入力の名前はそのまま
    assert "1.9GB" in models.warning("qwen2.5:3b")              # 変更時は必ずサイズを伝える


def test_ollama_members():
    import runtime
    names = ["ollama.exe", "lib/ollama/ggml-cpu.dll", "lib/ollama/cuda_v12/cublas.dll",
             "lib/ollama/cuda_v13/cublas.dll", "lib/ollama/vulkan/ggml-vulkan.dll"]
    assert [n for n in names if runtime.ollama_member(n, None)] == [names[0], names[1], names[4]]
    assert [n for n in names if runtime.ollama_member(n, "cuda_v13")] == [names[0], names[1], names[3], names[4]]


def test_quick_route():
    assert llm.quick_route("リンゴの写真探しておいて") == ("image_search", {"query": "リンゴ"})
    assert llm.quick_route("Excelを起動して") == ("launch_app", {"name": "Excel"})
    assert llm.quick_route("東京タワーについて調べてください") == ("search_summary", {"query": "東京タワー"})
    assert llm.quick_route("あなたの名前を教えて") is None          # 会話はAIに任せる
    assert llm.quick_route("ブラウザで天気を検索して") is None      # ブラウザ指定はAIに任せる


def test_streaming_chat(monkeypatch):
    import json

    class Resp:
        def __init__(self, lines):
            self.lines = lines

        def __enter__(self):
            return self

        def __exit__(self, *a):
            pass

        def raise_for_status(self):
            pass

        def iter_lines(self):
            return [json.dumps({"message": {"content": c}}).encode() for c in self.lines]

    got = []
    monkeypatch.setattr(llm, "on_delta", got.append)
    monkeypatch.setattr(llm.requests, "post", lambda *a, **k: Resp(["こんに", "ちは。", "元気です"]))
    assert llm._chat([], stream_out=True)["content"] == "こんにちは。元気です"
    assert got == ["こんに", "ちは。", "元気です"]               # 届いた分から順に渡す
    got.clear()
    monkeypatch.setattr(llm.requests, "post", lambda *a, **k: Resp(['{"name": "lock_screen"', "}"]))
    llm._chat([], stream_out=True)
    assert got == []                                           # 関数呼び出しのJSONは画面に出さない


def test_split_sentences():
    from tts import split_sentences
    assert split_sentences("はい。わかりました！次は？") == ["はい。", "わかりました！", "次は？"]
    assert split_sentences("短い、文。") == ["短い、文。"]


def test_threads(tmp_path, monkeypatch):
    import threads
    monkeypatch.setattr(threads, "DIR", str(tmp_path))
    t = threads.new()
    assert threads.listing() == []                      # まだ何も話していないスレッドは出さない
    threads.add(t, "user", "ポメラニアンの性格を調べて")
    threads.add(t, "ai", "明るい性格です。", sources="[1] A\nhttps://a")
    threads.add(t, "user", "寿命は？")
    threads.add(t, "ai", "12〜16年ほどです。")
    assert t["title"] == "ポメラニアンの性格を調べて"
    assert threads.history(t, 1) == [("寿命は？", "12〜16年ほどです。")]
    assert threads.load(t["id"])["messages"][1]["sources"] == "[1] A\nhttps://a"
    assert [x["id"] for x in threads.listing()] == [t["id"]]
    threads.delete(t["id"])
    assert threads.listing() == []


def test_space_ptt(monkeypatch):
    """スペースキー: 長押しで話す、短く押せば普通のスペースを入力し直す。"""
    import time
    import types
    import stopper
    events, sent, hooks = [], [], []
    fake = types.SimpleNamespace(KEY_DOWN="down", hook=lambda cb, suppress=False: hooks.append(cb) or (lambda: None),
                                 unhook=lambda h: None, send=lambda k: sent.append(k))
    monkeypatch.setitem(__import__("sys").modules, "keyboard", fake)
    stopper.register_ptt("space", lambda: events.append("start"), lambda: events.append("stop"))
    press = lambda t: hooks[-1](types.SimpleNamespace(name="space", event_type=t))  # noqa: E731
    assert press("down") is False and press("up") is False          # 短く押す: 一旦止めて…
    time.sleep(0.05)
    assert sent == ["space"] and events == []                         # 普通のスペースとして入力し直す
    assert press("down") is True and press("up") is True             # 入力し直した分は通す
    press("down")
    time.sleep(stopper.HOLD_SEC + 0.1)
    press("up")
    assert events == ["start", "stop"]                                # 長押し: 話す
    assert hooks[-1](types.SimpleNamespace(name="a", event_type="down")) is True  # 他のキーは通す


def test_transcribe_tuning():
    import numpy as np
    from voice import VoiceInput, SR
    calls = []

    class FakeModel:
        def __init__(self, text):
            self.text = text

        def transcribe(self, audio, **kw):
            calls.append(kw)
            return [type("S", (), {"text": self.text})()], None

    speech = (np.sin(np.arange(SR) / 5) * 0.05).astype(np.float32)   # 小さめの声 1秒
    assert VoiceInput.transcribe(FakeModel("メモ帳を開いて"), speech, True) == "メモ帳を開いて"
    assert calls[-1]["vad_filter"] is False                              # ボタンで録った音は切り落とさない
    assert VoiceInput.transcribe(FakeModel("ご視聴ありがとうございました"), speech, True) == ""  # 決まり文句は捨てる
    assert VoiceInput.transcribe(FakeModel("x"), np.zeros(SR, np.float32), True) == ""             # 無音
    assert VoiceInput.transcribe(FakeModel("x"), speech[:SR // 10], True) == ""                     # 短すぎる


def test_search_relevance():
    import json
    import websearch
    data = {"props": {"pageProps": {"initialProps": {"pageData": {"algos": [
        {"title": "<b>株式会社 才</b>（さい）", "url": "https://sai.example.jp/", "description": "Web制作の<b>株式会社才</b>"}]}}}}}
    page = f'<script id="__NEXT_DATA__" type="application/json">{json.dumps(data)}</script>'
    assert websearch.parse_yahoo_web(page) == [
        {"title": "株式会社 才（さい）", "url": "https://sai.example.jp/", "snippet": "Web制作の株式会社才"}]
    general = {"title": "株式会社とは？仕組みを解説", "url": "https://kaisha.example.jp/", "snippet": "株式会社の基本"}
    exact = {"title": "株式会社 才 会社概要", "url": "https://sai.example.jp/", "snippet": ""}
    assert websearch.is_relevant("株式会社才", exact) and not websearch.is_relevant("株式会社才", general)
    assert websearch.relevance("株式会社才", exact) > websearch.relevance("株式会社才", general)


def test_image_queries():
    import websearch
    assert websearch.image_queries("きれいなりんご") == ["りんご きれい", "りんご 高画質"]
    assert llm.quick_route("きれいなりんごの画像調べて") == ("image_search", {"query": "きれいなりんご"})


def test_wake_only_opens_chat():
    from wakeword import strip_wake_word
    assert strip_wake_word("ウェイク、ゼラス。") == ""
    assert strip_wake_word("ウェーク ゼラス、天気を教えて") == "天気を教えて"
    assert strip_wake_word("今日はいい天気") is None


def test_safe_only_and_chunks(monkeypatch):
    import websearch
    monkeypatch.setattr(websearch, "host_is_safe", lambda h: h != "bad.example")
    items = [{"url": "https://ok.example/a"}, {"url": "https://bad.example/x"}, {"url": "http://1.2.3.4/"}]
    assert websearch.safe_only(items, "url") == [items[0]]
    a = [{"full": f"https://ok.example/a{i}", "page": "https://ok.example"} for i in range(8)]
    b = [{"full": f"https://ok.example/b{i}", "page": "https://ok.example"} for i in range(8)]
    monkeypatch.setattr(websearch, "image_queries", lambda q: ["x", "y"])
    monkeypatch.setattr(websearch, "_images_one", lambda q, n: a if q == "x" else b)
    got = [it["full"].rsplit("/", 1)[1] for it in websearch.images("q", 10)]
    assert got == ["a0", "a1", "a2", "a3", "b0", "b1", "b2", "b3", "a4", "a5"]


def test_persona_parse():
    import persona
    assert persona.parse("関西弁で話して") == ("set", "関西弁で話して")
    assert persona.parse("これからは兄弟みたいに話して") == ("set", "兄弟みたいに話して")
    assert persona.parse("話し方を元に戻して") == ("reset", "")
    assert persona.parse("今日の天気を教えて") is None
    assert persona.apply("関西弁で話す", ("add", "もっと短く答えて")) == "関西弁で話す。もっと短く答えて"
    config.PERSONA = "関西弁で話す"
    assert "関西弁" in llm.system_prompt()
    config.PERSONA = ""


def test_voice_danger_needs_confirm(tmp_path, monkeypatch):
    monkeypatch.setattr(config, "DESKTOP", str(tmp_path))
    (tmp_path / "memo.txt").write_text("x", encoding="utf-8")
    monkeypatch.setattr(tools, "REQUIRE_CONFIRM", True)
    res = tools.run_tool("delete_item", {"name": "memo.txt"})
    body, confirm = llm.take_confirm(res)
    assert confirm["tool"] == "delete_item" and "はい" in body
    assert (tmp_path / "memo.txt").exists()  # 確認前なので消えていない
    assert tools.CONFIRM not in tools.run_tool("create_folder", {"name": "a"})  # 危なくない操作はそのまま
    monkeypatch.setattr(tools, "REQUIRE_CONFIRM", False)
    assert tools.CONFIRM not in tools.run_tool("shutdown_pc", {"cancel": True})


def test_open_url_rejects_bad(monkeypatch):
    import websearch
    monkeypatch.setattr(websearch, "host_is_safe", lambda h: h != "bad.example")
    opened = []
    monkeypatch.setattr(tools, "_browser_path", lambda: None)
    monkeypatch.setattr(tools.webbrowser, "open", lambda u: opened.append(u))
    import browser
    monkeypatch.setattr(browser, "enabled", lambda: False)
    assert "開きませんでした" in tools.run_tool("open_url", {"url": "https://bad.example/x"})
    assert "開きませんでした" in tools.run_tool("open_url", {"url": "http://1.2.3.4/"})
    tools.run_tool("open_url", {"url": "https://ok.example/"})
    assert opened == ["https://ok.example/"]


def test_reading_fix():
    import reading
    assert reading.fix("確認を行っています") == "確認をおこなっています"
    assert reading.fix("東京へ行っています") == "東京へ行っています"
    assert reading.fix("会議を行います") == "会議をおこないます"
    assert reading.fix("流行っている") == "流行っている"


def test_off_track_reply_is_retried(monkeypatch):
    import json as _json
    replies = iter(["What is the weather like tomorrow in Tokyo?", "明日の東京は晴れの予想です。"])

    class Resp:
        def __init__(self, text):
            self.lines = [_json.dumps({"message": {"content": ch}}).encode() for ch in text]

        def __enter__(self):
            return self

        def __exit__(self, *a):
            return False

        def raise_for_status(self):
            pass

        def iter_lines(self):
            return iter(self.lines)

    sent = []
    monkeypatch.setattr(llm.requests, "post", lambda *a, **k: (sent.append(k["json"]), Resp(next(replies)))[1])
    shown = []
    monkeypatch.setattr(llm, "on_delta", shown.append)
    monkeypatch.setattr(config, "QUICK_ROUTE", False)
    reply = llm.ask("明日の東京の天気はどう？", history=[])
    assert reply == "明日の東京は晴れの予想です。"
    assert "".join(shown) == reply  # 英語の書き出しは画面にも読み上げにも出ていない
    assert len(sent) == 2 and any(m["role"] == "system" and "翻訳" in m["content"] for m in sent[1]["messages"])


def test_speaker_plays_continuously_and_stops(monkeypatch):
    import sys
    import time as _t
    import types
    import numpy as np
    import tts
    written = []

    class Out:
        def __init__(self, samplerate, channels, dtype, latency):
            self.channels, self.latency, self.active = channels, 0.01, False

        def start(self):
            self.active = True

        def write(self, data):
            written.append(data.copy())
            _t.sleep(0.005)

        def abort(self):
            self.active = False

        def close(self):
            pass

    monkeypatch.setitem(sys.modules, "sounddevice", types.SimpleNamespace(OutputStream=Out, stop=lambda: None))
    sp = tts.Speaker()
    a = (np.arange(24000, dtype=np.int16).reshape(-1, 1), 24000)  # 1秒
    b = (np.arange(12000, dtype=np.int16).reshape(-1, 1), 24000)  # 0.5秒
    sp._audio.put((sp._gen, a, "一つ目"))
    sp._audio.put((sp._gen, b, "二つ目"))
    for _ in range(200):
        if sum(len(w) for w in written) >= 36000:
            break
        _t.sleep(0.01)
    got = np.concatenate(written)[:, 0]
    assert len(got) == 36000 and (got[:24000] == a[0][:, 0]).all() and (got[24000:] == b[0][:, 0]).all()
    # 途中で止めると、すぐに書き込みが止まる
    written.clear()
    long = (np.zeros((24000 * 10, 1), np.int16), 24000)
    sp._audio.put((sp._gen, long, "長い"))
    _t.sleep(0.1)
    sp.stop()
    _t.sleep(0.2)
    assert sum(len(w) for w in written) < 24000 * 10


def test_memory_and_context(monkeypatch):
    import memory
    assert memory.parse("私の誕生日は5月3日だって覚えておいて") == ("remember", "私の誕生日は5月3日")
    assert memory.parse("会社名は株式会社才、覚えといて") == ("remember", "会社名は株式会社才")
    assert memory.parse("誕生日のことは忘れて") == ("forget", "誕生日")
    assert memory.parse("明日の予定を教えて") is None
    assert memory.apply(["誕生日は5月3日", "好きな色は青"], ("forget", "誕生日")) == ["好きな色は青"]
    sent = []
    monkeypatch.setattr(llm, "_chat", lambda m, use_tools=True, **kw: (sent.append(m), {"content": "はい"})[1])
    monkeypatch.setattr(config, "QUICK_ROUTE", False)
    llm.ask("それは何色？", history=[("好きな色を当てて", "青ですか？")], summary="ユーザーは色の話をしている",
            memory=["好きな色は青"])
    msgs = sent[0]
    assert "好きな色は青" in msgs[0]["content"]  # 覚えていることは指示文に
    assert any(m["role"] == "system" and "色の話" in m["content"] for m in msgs[1:])  # 要約
    assert msgs[-1]["content"] == "それは何色？" and msgs[-2]["content"] == "青ですか？"  # 直前のやりとり


def test_export_text():
    import threads
    t = threads.new()
    t["messages"] = [{"role": "user", "text": "こんにちは", "time": 0}, {"role": "ai", "text": "やあ", "time": 0}]
    out = threads.export_text(t, "Zealous")
    assert "**あなた**" in out and "こんにちは" in out and "**Zealous**" in out


def test_find_and_open_files(tmp_path, monkeypatch):
    monkeypatch.setattr(config, "DESKTOP", str(tmp_path))
    monkeypatch.setitem(tools._KNOWN, "デスクトップ", str(tmp_path))
    monkeypatch.setattr(tools, "_search_roots", lambda path="": [str(tmp_path)])
    (tmp_path / "テンプレ.txt").write_text("x", encoding="utf-8")
    (tmp_path / "sub").mkdir()
    (tmp_path / "sub" / "テンプレ_古い.txt").write_text("y", encoding="utf-8")
    res = tools.run_tool("find_files", {"name": "テンプレ"})
    assert "2 件" in res and str(tmp_path / "テンプレ.txt") in res
    opened = []
    monkeypatch.setattr(tools.subprocess, "Popen", lambda args, **k: opened.append(args))
    if hasattr(tools.os, "startfile"):
        monkeypatch.setattr(tools.os, "startfile", lambda p: opened.append(p))
    assert "開きました" in tools.run_tool("open_file", {"name": "テンプレ"})
    assert str(tmp_path / "テンプレ.txt") in str(opened[0])
    # すでにあるファイルを、空の中身で上書きしない
    assert "すでにあります" in tools.run_tool("create_text_file", {"name": "テンプレ"})
    assert (tmp_path / "テンプレ.txt").read_text(encoding="utf-8") == "x"


def test_claim_without_tool_is_retried(monkeypatch):
    calls = iter([{"content": "ファイルを開きました。"}, {"content": "", "tool_calls": [
        {"function": {"name": "find_files", "arguments": {"name": "テンプレ"}}}]}])
    monkeypatch.setattr(llm, "_chat", lambda m, use_tools=True, **kw: next(calls))
    monkeypatch.setattr(tools, "find_paths", lambda name, path="", limit=10: ["C:/Users/u/Desktop/テンプレ.txt"])
    monkeypatch.setattr(config, "QUICK_ROUTE", False)
    monkeypatch.setattr(config, "SUMMARIZE_AFTER_TOOL", False)
    assert "見つかりました" in llm.ask("さっきのファイルを探してきて", history=[])


def test_web_routes_and_invented_urls(monkeypatch):
    opened = []
    monkeypatch.setattr(tools, "open_in_browser", lambda u: opened.append(u))
    assert llm.quick_route("ブラウザからYouTube開いて、株式会社才の動画を流して") == ("play_video", {"query": "株式会社才"})
    assert llm.quick_route("Amazonで充電器を検索して") == ("open_site", {"site": "Amazon", "query": "充電器"})
    assert "推測" in tools.run_tool("open_url", {"url": "https://www.youtube.com/watch?v=example_video_id"})
    assert opened == []
    tools.run_tool("open_site", {"site": "楽天", "query": "コーヒー豆"})
    assert opened[-1].startswith("https://search.rakuten.co.jp/search/mall/")


def test_scroll_follow_ups(monkeypatch):
    monkeypatch.setattr(tools, "run_tool", lambda n, a: "「ページ」をスクロールしました")
    monkeypatch.setitem(llm._last_scroll, "direction", None)
    steps = []
    for t in ["下にスクロールして", "もっと下に", "もっと下に", "少し戻って"]:
        r = llm.quick_route(t)
        llm._run(*r)
        steps.append((r[1]["direction"], r[1]["amount"]))
    assert steps == [("down", 5), ("down", 5), ("down", 5), ("up", 2)]


def test_play_list(monkeypatch):
    opened = []
    monkeypatch.setattr(tools, "open_in_browser", lambda u: opened.append(u))
    monkeypatch.setattr(tools, "youtube_search", lambda q, order="relevance", limit=10:
                        [(f"id{i:09d}", f"曲{i}") for i in range(limit)])
    r = llm.quick_route("YOASOBIの曲を再生回数順に流して")
    assert r == ("play_video", {"query": "YOASOBI", "order": "views", "count": 10})
    res = tools.run_tool(*r)
    assert "再生回数の多い順に10本" in res
    assert opened[-1].startswith("https://www.youtube.com/watch_videos?video_ids=id000000000,id000000001")


def test_news_and_followup_search(monkeypatch):
    assert llm.quick_route("AIに関する最新のビッグニュースいくつかしょうかいして") == \
        ("search_summary", {"query": "AIに関する最新のビッグニュース"})
    llm._history.clear()
    llm._history.append(("AIに関する最新のビッグニュースいくつかしょうかいして", "できません"))
    assert llm.quick_route("じゃあそれ調べて") == ("search_summary", {"query": "AIに関する最新のビッグニュース"})
    text = '調べてみましょう。\n\n[Tools]\n{"type": "function", "function": {web_search "AIの最新ニュース"}}'
    assert llm.parse_tool_calls({"content": text}) == [("web_search", {"query": "AIの最新ニュース"})]
    assert llm.visible(text) == "調べてみましょう。"
    # 「最新情報は提供できません」と断ったら、実際に調べる
    monkeypatch.setattr(config, "QUICK_ROUTE", False)
    monkeypatch.setattr(llm, "_chat", lambda m, use_tools=True, **kw: {"content": "私は現在の情報を提供する能力はありません。"})
    monkeypatch.setattr(llm, "_run", lambda n, a: f"{n}:{a['query']}")
    assert llm.ask("元気？最新情報ある？", history=[]).startswith("search_summary")


def test_weather_routes(monkeypatch):
    import weather
    assert llm.quick_route("明日の東京の天気は") == ("get_weather", {"place": "東京", "day": "明日"})
    assert llm.quick_route("明日大阪は雨降る") == ("get_weather", {"place": "大阪", "day": "明日"})
    llm._history.clear()
    llm._history.append(("明日の東京の天気は？", "できません"))
    assert llm.quick_route("ネットで調べてよ") == ("get_weather", {"place": "東京", "day": "明日"})
    monkeypatch.setattr(weather, "forecast", lambda lat, lon, days=3: {
        "time": ["d0", "d1", "d2"], "weather_code": [0, 63, 3], "temperature_2m_max": [20, 18.4, 22],
        "temperature_2m_min": [10, 12.6, 14], "precipitation_probability_max": [0, 80, 10]})
    assert tools.run_tool("get_weather", {"place": "東京", "day": "明日"}) == \
        "明日の東京は雨、最高18度・最低13度、降水確率80%の予報です。"


def test_web_page_routes(monkeypatch):
    import browser
    monkeypatch.setattr(browser, "enabled", lambda: True)
    monkeypatch.setattr(browser, "active", lambda: True)
    assert llm.quick_route("検索欄にねこと入力して検索して") == \
        ("web_type", {"text": "ねこ", "submit": True, "field": "検索"})
    assert llm.quick_route("ログインを押して") == ("web_click", {"target": "ログイン"})
    assert llm.quick_route("エンターを押して") == ("web_key", {"keys": "エンター"})
    assert llm.quick_route("前のページに戻って") == ("web_page", {"action": "back"})
    assert llm.quick_route("10秒飛ばして") == ("web_video", {"action": "forward", "seconds": 10})
    assert llm.quick_route("このページを要約して") == ("web_read", {"question": ""})
    monkeypatch.setattr(browser, "active", lambda: False)  # ブラウザが手前に無ければ、ふつうの扱い
    assert llm.quick_route("ログインを押して") is None


def test_full_context_compress(monkeypatch):
    import threads
    t = threads.new()
    for i in range(30):
        threads.add(t, "user", f"質問{i}" + "あ" * 200)
        threads.add(t, "ai", f"答え{i}" + "い" * 400)
    pairs, summary, base = threads.full_history(t)
    assert len(pairs) == 30 and summary == "" and base == 0 and len(pairs[0][1]) > 400  # 省略しない
    monkeypatch.setattr(config, "FULL_CONTEXT_LIMIT", 8000)
    got = []
    monkeypatch.setattr(llm, "on_summary", lambda tid, text, upto: got.append((tid, text, upto)))
    monkeypatch.setattr(llm, "_chat", lambda m, use_tools=True, **kw: {"content": "要約です"})
    keep, new = llm._fit(pairs, "", t["id"], 0)
    assert new == "要約です" and sum(len(u) + len(a) for u, a in keep) <= 4800
    assert got == [(t["id"], "要約です", 30 - len(keep))]
    t.update(summary="要約です", summary_upto=got[0][2], summary_full=True)
    assert threads.full_history(t)[1:] == ("要約です", got[0][2])
    monkeypatch.setattr(config, "FULL_CONTEXT_COMPRESS", False)  # 圧縮しない: 古い部分は参照しない
    assert llm._fit(pairs, "", t["id"], 0)[1] == ""
    assert llm.num_ctx([{"content": "あ" * 100}]) == 4096 and llm.num_ctx([{"content": "あ" * 12000}]) == 16384


def test_app_name_variants_and_choice(monkeypatch):
    import apps
    found = [{"Name": n, "AppID": n} for n in
             ("Mozilla Thunderbird", "Discord", "Microsoft Edge", "Zoom Workplace", "Zoom Rooms")]
    assert [a["Name"] for a in apps.matches("サンダーバード", found)] == ["Mozilla Thunderbird"]
    assert [a["Name"] for a in apps.matches("でぃすこーど", found)] == ["Discord"]
    assert apps.matches("請求書", found, fuzzy=False) == []
    monkeypatch.setattr(apps, "start_apps", lambda: found)
    monkeypatch.setattr(apps.subprocess, "Popen", lambda *a, **k: None)
    llm._choice = None
    reply = llm._run("launch_app", {"name": "ズーム"})  # 「Zoom」が2つ: どちらか聞く
    assert "2つ" in reply and "1. Zoom Rooms" in reply and tools.CHOICE not in reply and llm._choice
    monkeypatch.setattr(config, "QUICK_ROUTE", False)
    assert llm._ask("2番で") == "Zoom Workplace を起動しました"
    llm._run("launch_app", {"name": "ズーム"})
    assert llm._ask("やめて") == "わかりました。やめておきます。" and llm._choice is None
    assert llm._pick("ワークプレイスの方", {"options": ["a", "b"], "labels": ["Zoom Rooms", "Zoom ワークプレイス"]}) == 1


def test_streamed_text_is_retracted_when_tool_runs(monkeypatch):
    monkeypatch.setattr(config, "QUICK_ROUTE", False)
    sent = []
    monkeypatch.setattr(llm, "on_retract", lambda: sent.append("retract"))

    def fake_chat(messages, use_tools=True, **kw):
        llm._streamed = True  # 「サンダーバードを開きます」と書き始めてから、関数を呼んだ
        return {"content": "サンダーバードを開きます", "tool_calls": [
            {"function": {"name": "launch_app", "arguments": {"name": "サンダーバード"}}}]}
    monkeypatch.setattr(llm, "_chat", fake_chat)
    monkeypatch.setattr(llm, "_run", lambda n, a: "「サンダーバード」というアプリが見つかりませんでした")
    monkeypatch.setattr(config, "SUMMARIZE_AFTER_TOOL", False)
    assert llm._ask("サンダーバード開いて") == "「サンダーバード」というアプリが見つかりませんでした"
    assert sent == ["retract"]


def test_mail_routes_and_imap(monkeypatch):
    import mail
    assert llm.quick_route("メールのアプリを開いて、迷惑メールフォルダのメールを全て既読にして") == \
        ("mail_mark_read", {"folder": "迷惑メール"})
    assert llm.quick_route("メールの仕事フォルダの名前を才コンサルに変えて") == \
        ("mail_rename_folder", {"folder": "仕事", "new_name": "才コンサル"})
    assert mail._utf7_decode(mail._utf7_encode("迷惑メール&A")) == "迷惑メール&A"

    class FakeImap:
        flags = {b"1": False, b"2": False, b"3": True}
        renamed = []

        def list(self):
            return "OK", [b'(\\HasNoChildren) "/" "INBOX"',
                          b'(\\HasNoChildren \\Junk) "/" "[Gmail]/' + mail._utf7_encode("迷惑メール").encode() + b'"',
                          b'(\\HasNoChildren) "/" "' + mail._utf7_encode("仕事").encode() + b'"']

        def select(self, name):
            self.selected = name

        def search(self, charset, crit):
            return "OK", [b" ".join(k for k, seen in self.flags.items() if not seen)]

        def store(self, ids, op, flag):
            for i in ids.split(","):
                self.flags[i.encode()] = True

        def rename(self, old, new):
            self.renamed.append((old, new))
            return "OK", [b""]

    fake = FakeImap()

    class Ctx:
        def __enter__(self):
            return fake

        def __exit__(self, *a):
            pass
    monkeypatch.setattr(mail, "_Imap", Ctx)
    monkeypatch.setattr(config, "MAIL_IMAP_HOST", "imap.example.com")
    monkeypatch.setattr(config, "MAIL_IMAP_USER", "me@example.com")
    assert tools.run_tool("mail_mark_read", {"folder": "迷惑メール"}) == "「迷惑メール」の未読メール2通を、すべて既読にしました"
    assert fake.selected == '"[Gmail]/' + mail._utf7_encode("迷惑メール") + '"'
    monkeypatch.setattr(llm, "_chat", lambda m, use_tools=True, **kw: {"content": "仕事\n私事"})
    assert tools.run_tool("mail_rename_folder", {"folder": "しごと", "new_name": "才コンサル"}).endswith("「才コンサル」に変えました")
    assert fake.renamed == [('"' + mail._utf7_encode("仕事") + '"', '"' + mail._utf7_encode("才コンサル") + '"')]


def test_find_in_files_by_content(monkeypatch, tmp_path):
    (tmp_path / "a.txt").write_text("来月の旅行の計画。京都へ行く", encoding="utf-8")
    (tmp_path / "b.txt").write_bytes("買い物リスト".encode("cp932"))
    (tmp_path / "c.md").write_text("旅行 リョコウ メモ", encoding="utf-8")
    monkeypatch.setattr(config, "DESKTOP", str(tmp_path))
    monkeypatch.setattr(tools, "_KNOWN", dict(tools._KNOWN, デスクトップ=str(tmp_path)))
    opened = []
    monkeypatch.setattr(tools, "open_file", lambda p, path="": opened.append(p) or "開きました")
    assert [os.path.basename(p) for p, _, _ in tools.search_contents("買い物", "デスクトップ")] == ["b.txt"]
    assert tools.run_tool("find_in_files", {"text": "京都", "path": "デスクトップ", "kind": "メモ帳", "open": True}) == "開きました"
    assert tools.run_tool("find_in_files", {"text": "りょこう", "path": "デスクトップ", "open": True}) == "開きました"
    assert opened[-1].endswith("c.md")  # カタカナで書かれていても見つかる
    reply = tools.run_tool("find_in_files", {"text": "旅行", "path": "デスクトップ", "open": True})
    assert "2件" in reply and tools.CHOICE in reply  # 2つ: どれを開くか聞く


def test_voice_names_ask_spelling(monkeypatch, tmp_path):
    monkeypatch.setattr(llm, "_chat", lambda m, use_tools=True, **kw: {"content": "1. 才コンサル\n2. Sai Consul"})
    assert tools.spell_options("さいこんさる.txt") == ["さいこんさる.txt", "サイコンサル.txt", "才コンサル.txt", "Sai Consul.txt"]
    monkeypatch.setattr(tools, "REQUIRE_CONFIRM", True)
    monkeypatch.setattr(config, "DESKTOP", str(tmp_path))
    llm._choice = None
    reply = llm._run("create_folder", {"name": "さいこんさる"})
    assert "どの書き方" in reply and llm._choice
    monkeypatch.setattr(config, "QUICK_ROUTE", False)
    llm._ask("3番")
    assert (tmp_path / "才コンサル").is_dir()


def test_technical_mode(monkeypatch, tmp_path):
    assert llm.looks_technical("nginxのリバースプロキシの設定方法は？")
    assert llm.looks_technical("def f():\n    return 1\nこれ直して")
    assert not llm.looks_technical("明日の天気は？") and not llm.looks_technical("コマンドプロンプトを開いて")
    src = tmp_path / "app.py"
    src.write_text("print('hi'", encoding="utf-8")
    seen = []

    def fake_chat(messages, use_tools=True, **kw):
        seen.append((messages[0]["content"][:30], [m["role"] for m in messages]))
        if messages[-1]["role"] == "user":  # まずファイルを読む
            return {"role": "assistant", "content": "", "tool_calls": [{"function": {"name": "read_text_file", "arguments": {"name": str(src)}}}]}
        assert "print('hi'" in messages[-1]["content"]
        return {"content": "括弧が閉じていません。\n```python\nprint('hi')\n```"}
    monkeypatch.setattr(llm, "_chat", fake_chat)
    reply = llm.ask(f"{src} を添削して", history=[], mode="tech", full=True)
    assert reply.endswith("```") and "```python" in reply  # コードはそのまま残す
    assert "テクニカルモード" in seen[0][0] or "エンジニア" in seen[0][0]
    llm._tech = False


def test_speech_skips_code():
    import tts
    f = tts.CodeFilter()
    out = f.feed("直しました。\n``") + f.feed("`python\nprint(1)\n``") + f.feed("`\n以上です。")
    assert "print" not in out and out.startswith("直しました。") and out.endswith("以上です。")


def test_office_files(monkeypatch, tmp_path):
    import officedoc
    text = ("# 提案書\n概要です\n\n## 背景\n- 人手不足\n- **自動化**の需要\n> メモ\n\n"
            "## 売上\n| 年度 | 売上(円) | 率 |\n|---|---|---|\n| 2026 | 1,200 | 12% |\n| 2027 | 3,500 | 18% |\n\n"
            "## 手順\n1. 整理\n2. 試験\n3. 導入\n")
    doc = officedoc.parse(text)
    assert doc["title"] == "提案書" and doc["subtitle"] == "概要です" and len(doc["sections"]) == 3
    assert doc["sections"][1]["blocks"][0][0] == "table" and doc["sections"][0]["notes"] == ["メモ"]
    monkeypatch.setattr(llm, "write_document", lambda kind, request, source="": text)
    monkeypatch.setattr(config, "DESKTOP", str(tmp_path))
    for kind, unit in (("pptx", "スライド4枚"), ("docx", "3章"), ("xlsx", "1シート")):
        reply = tools.run_tool("make_office_file", {"kind": kind, "request": "提案書を作って"})
        assert unit in reply and (tmp_path / f"提案書.{kind}").exists(), reply
    assert "提案書 (2).pptx" in tools.run_tool("make_office_file", {"kind": "パワポ", "request": "x"})  # 上書きしない
    from openpyxl import load_workbook
    ws = load_workbook(tmp_path / "提案書.xlsx").active
    assert ws["B4"].value == 1200 and ws["C4"].value == 0.12 and ws["A4"].number_format == "0"
    assert llm.quick_route("新サービスの提案資料をパワポで作って")[0] == "make_office_file"


def test_document_research_review_and_images(monkeypatch, tmp_path):
    import officedoc
    import websearch
    monkeypatch.setattr(websearch, "search", lambda q, n=5: [{"title": "白書", "url": "https://example.go.jp/a",
                                                              "snippet": "導入率は30%"}])
    monkeypatch.setattr(websearch, "read_pages", lambda r, count=3, limit=1500: ["本文"])
    calls = []

    def fake(messages, model, temperature):
        calls.append(messages[0]["content"])
        if "校正" in messages[0]["content"]:
            return "# 題名\n## 背景\n- 導入率は30%\n## まとめ\n- 小さく始める"
        return "# 題名\n## 背景\n- 導入率わ30%\n- 導入率わ30%\n## まとめ\n- 小さく始める\n!画像: office"
    monkeypatch.setattr(llm, "_doc_chat", fake)
    text = llm.write_document("pptx", "生成AIの導入についての提案資料をパワポで作って")
    assert "【調べた情報】\n" in calls[0] and "導入率は30%" in calls[0]  # 調べた内容だけを根拠にする
    assert "校正" in calls[1] and text.endswith("## 出典\n- 白書 https://example.go.jp/a")
    assert llm.doc_topic("生成AIの導入についての提案資料をパワポで作って") == "生成AIの導入についての提案資料"
    calls.clear()
    llm.write_document("docx", "今の内容をWordにまとめて")  # 会話の内容から作るときは調べない
    assert "【調べた情報】\n" not in calls[0]
    # 重複は1つに、出典は最後の一覧に
    doc = officedoc.parse("# T\n## A\n- 同じ\n- 同じ。\n## B\n- 別\n## 出典\n- 白書 https://x")
    assert [b for _, b in doc["sections"][0]["blocks"]] == [(0, "同じ")] and doc["sources"] == ["白書 https://x"]
    # 写真は許可の明記されたものだけ(取れなければその行を使わない)
    monkeypatch.setattr(officedoc, "commons_image", lambda q, folder: None)
    doc = officedoc.parse("# T\n## A\n- a\n!画像: cat")
    officedoc.attach_images(doc, str(tmp_path))
    assert all(k != "image" for k, _ in doc["sections"][0]["blocks"])


def test_attachments_read_and_answer(monkeypatch, tmp_path):
    import attachments
    import officedoc
    (tmp_path / "proj").mkdir()
    (tmp_path / "proj" / "README.md").write_text("# 在庫管理ツール\n商品の在庫を数える", encoding="utf-8")
    (tmp_path / "proj" / "main.py").write_text("def count(items):\n    return len(items)\n", encoding="utf-8")
    (tmp_path / "sales.csv").write_text("月,売上\n4月,100\n5月,300\n", encoding="utf-8")
    doc = officedoc.parse("# 報告\n## 結果\n- 売上は増えた")
    officedoc.make_docx(doc, str(tmp_path / "rep.docx"))
    officedoc.make_xlsx(officedoc.parse("# 表\n## 売上\n| 月 | 売上 |\n|---|---|\n| 4月 | 100 |\n| 5月 | 300 |"),
                        str(tmp_path / "t.xlsx"))
    text = attachments.material([str(tmp_path / "proj"), str(tmp_path / "sales.csv"), str(tmp_path / "rep.docx"),
                                 str(tmp_path / "t.xlsx")])
    assert "【フォルダ: proj/】" in text and "README.md" in text and "在庫管理ツール" in text and "def count" in text
    assert "2行 × 2列" in text and "合計 400" in text and "売上は増えた" in text and "シート「売上」" in text
    seen = []
    monkeypatch.setattr(llm, "_chat_on_track", lambda m, q, use_tools=False: seen.append(m) or {"content": "合計は400です"})
    monkeypatch.setattr(config, "QUICK_ROUTE", True)
    assert llm.ask("売上の合計は？", history=[], attachments=[str(tmp_path / "sales.csv")]) == "合計は400です"
    assert "合計 400" in seen[0][1]["content"]  # 添付の中身をもとに答える
    llm.attach([])


def test_tech_model_only_in_tech_mode(monkeypatch):
    monkeypatch.setattr(config, "TECH_MODEL", "maternion/lfm2:8b-a1b")
    used, released = [], []
    monkeypatch.setattr(llm, "_doc_chat", lambda m, model, t: used.append(model) or "# T\n## A\n- a")
    monkeypatch.setattr(config, "OFFICE_RESEARCH", False)
    monkeypatch.setattr(config, "OFFICE_REVIEW", False)
    llm._tech = False
    llm.write_document("docx", "報告書を作って")
    assert used == [config.OLLAMA_MODEL]  # 通常モードでは、テクニカルモード用のAIを使わない
    llm._tech = True
    llm.write_document("docx", "報告書を作って")
    assert used[-1] == "maternion/lfm2:8b-a1b"
    monkeypatch.setattr(llm, "release_tech_model", lambda: released.append(1))
    monkeypatch.setattr(llm, "_ask", lambda text: "ok")
    llm.ask("こんにちは", history=[], mode="")  # テクニカルモードを抜けたら、すぐメモリから外す
    assert released == [1] and llm._tech is False


def test_office_interview_before_making(monkeypatch, tmp_path):
    made = []
    monkeypatch.setattr(tools, "make_office_file", lambda **kw: made.append(kw) or "作りました")
    tools.TOOLS["make_office_file"] = (tools.make_office_file, tools.TOOLS["make_office_file"][1])
    monkeypatch.setattr(config, "QUICK_ROUTE", True)
    llm._office = None
    llm.attach([])
    q = llm.ask("新サービスの提案資料をパワポで作って", history=[])
    assert "「新サービスの提案資料」" in q and "1. インターネットで調べて作る" in q and not made
    q = llm.ask("3番、料金は月額3000円、サポート付き", history=[])
    assert "デザインの希望" in q  # その場で内容も言ったので、内容の質問は飛ばす
    assert llm.ask("緑系で明朝体、5枚くらい", history=[]) == "作りました"
    kw = made[0]
    assert kw["kind"] == "pptx" and kw["source"] == "user" and kw["design"] == "緑系で明朝体、5枚くらい"
    assert "料金は月額3000円" in kw["request"] and "【デザインの希望】" in kw["request"]
    # 内容がはっきりしないときは、まず何の資料かを聞く。やめることもできる
    q = llm.ask("ワードで資料作って", history=[])
    assert "どんな内容" in q
    assert llm.ask("やめて", history=[]) == "わかりました。資料づくりはやめておきます。" and llm._office is None
    # ネットで調べる・おまかせ
    llm.ask("会社紹介をエクセルで作って", history=[])
    llm.ask("ネットで調べて", history=[])
    llm.ask("おまかせ", history=[])
    assert made[-1]["source"] == "web" and made[-1]["design"] == "" and made[-1]["kind"] == "xlsx"
    import officedoc
    assert officedoc.theme_of("緑系で") == "green" and officedoc.theme_of("おまかせ") is None


def test_zip_attachment(tmp_path):
    import zipfile
    import attachments
    p = tmp_path / "proj.zip"
    with zipfile.ZipFile(p, "w") as z:
        z.writestr("proj/README.md", "# 在庫管理\n商品の在庫を数える")
        z.writestr("proj/data.csv", "月,売上\n4月,100\n5月,300\n")
        z.writestr("../evil.txt", "外へ出るパス")
        z.writestr("proj/big.txt", "あ" * 10)
    text = attachments.material([str(p)])
    assert "【ZIP: proj.zip】" in text and "在庫管理" in text and "合計 400" in text
    assert "外へ出るパス" not in text and "../evil.txt" in text  # 一覧には出すが、外へ出るパスは取り出さない
    assert not (tmp_path.parent / "evil.txt").exists()


def test_office_interview_flexible(monkeypatch):
    made = []
    monkeypatch.setattr(tools, "make_office_file", lambda **kw: made.append(kw) or "作りました")
    tools.TOOLS["make_office_file"] = (tools.make_office_file, tools.TOOLS["make_office_file"][1])
    monkeypatch.setattr(config, "QUICK_ROUTE", True)
    llm._office = None
    llm.attach([])
    # 形式も内容も言わない: まず両方を聞く。答えに形式・材料・内容が全部入っていれば、まとめて受け取る
    q = llm.ask("資料つくって", history=[])
    assert "どの形式" in q and "どんな内容" in q
    q = llm.ask("インターネットで株式会社才についてまとめてパワポにして", history=[])
    assert "デザイン" in q and "「株式会社才」" in q
    llm.ask("おまかせ", history=[])
    assert made[-1]["kind"] == "pptx" and made[-1]["source"] == "web" and made[-1]["request"].startswith("株式会社才の")
    # AIが勝手に考えた依頼文ではなく、ユーザーの言葉で聞き取りを始める
    llm._current_text = "資料を作成して"
    q = llm._run("make_office_file", {"kind": "docx", "request": "新プロジェクトの概要"})
    assert "新プロジェクト" not in q and "どんな内容" in q and llm._office["kind"] == "docx"
    # 途中で形式を変えても受け取る
    q = llm.ask("やっぱりエクセルで、店舗ごとの売上", history=[])
    assert llm._office["kind"] == "xlsx" and "中身はどうしますか" in q
    llm.ask("やめて", history=[])


def test_office_plan_to_slides(tmp_path):
    import json
    import llm
    import officedoc
    plan = {"title": "伝わる資料で成果を変える", "subtitle": "サービスのご紹介",
            "slides": [{"layout": "cards", "title": "特徴", "message": "3つの強みがあります。",
                        "items": [{"head": "速い", "text": "最短3日で納品"}, {"head": "安心", "text": "何度でも修正"}]},
                       {"layout": "steps", "title": "流れ", "message": "相談から納品まで。",
                        "items": [{"head": "相談", "text": "希望を聞く"}, {"head": "制作", "text": "形にする"}]},
                       {"layout": "compare", "title": "比較", "message": "違いは明らか。",
                        "items": [{"head": "自作", "text": "時間がかかる／ばらつく"}, {"head": "依頼", "text": "早い／整う"}]},
                       {"layout": "stats", "title": "数字", "message": "実績です。",
                        "items": [{"head": "365日", "text": "サポート"}, {"head": "3営業日", "text": "納品"}]}],
            "closing": {"title": "次の一歩", "points": ["まず相談", "気軽にどうぞ"]}}
    doc = officedoc.parse(llm._plan_md("pptx", plan))
    assert doc["title"] == "伝わる資料で成果を変える"
    assert doc["sections"][0]["lead"] == "3つの強みがあります。"
    assert officedoc._split_head(doc["sections"][0]["blocks"][0][1][1]) == ("速い", "最短3日で納品")
    assert officedoc.make_pptx(doc, str(tmp_path / "a.pptx")) >= 7  # 表紙・目次・4枚・まとめ
    d = llm._plan_md("docx", {"title": "T", "summary": "S", "key_points": ["要点1"],
                              "sections": [{"heading": "A", "paragraphs": ["本文"]}]})
    assert officedoc.parse(d)["key_points"] == ["要点1"]


def test_office_answers_and_titles():
    import officedoc
    assert officedoc.generic_title("資料") and officedoc.generic_title("") and not officedoc.generic_title("株式会社才")
    doc = officedoc.parse("## スライドの題名\nスライド1: 会社紹介\n■ 強み\n## 出典\n- a https://x.jp/?srsltid=abc")
    assert doc["sections"][0]["title"] == "会社紹介" and doc["sources"] == ["a https://x.jp/"]


def test_tech_model_fits_memory(monkeypatch):
    import config
    import llm
    import models
    monkeypatch.setattr(config, "OLLAMA_MODEL", "qwen2.5:1.5b")
    monkeypatch.setattr(config, "TECH_MODEL", "auto")
    monkeypatch.setattr(llm, "_tech_note", "")
    monkeypatch.setattr(models, "memory_gb", lambda: (8.0, 3.0))
    assert llm.tech_model() == "qwen2.5:1.5b"  # メモリ8GB: いつものAIのまま
    monkeypatch.setattr(models, "memory_gb", lambda: (16.0, 8.0))
    assert llm.tech_model() == "granite4.2:3b-q4_K_M"
    monkeypatch.setattr(config, "TECH_MODEL", "maternion/lfm2:8b-a1b")
    monkeypatch.setattr(models, "memory_gb", lambda: (16.0, 3.0))
    assert llm.tech_model() == "qwen2.5:1.5b" and "メモリが少ない" in llm._tech_note  # 空きが足りない
    monkeypatch.setattr(models, "memory_gb", lambda: (8.0, 6.0))
    assert llm.num_ctx([{"content": "あ" * 30000}]) == 8192  # メモリが少ないPCは読む量の上限を下げる
